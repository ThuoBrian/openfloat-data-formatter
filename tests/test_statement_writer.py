"""Tests for the Statement Report workbook — writer.write_statement_workbook.

Every report here is built from a synthetic statement (`make_statement_workbook`).
The real exports in sample_report_output/ hold staff names and phone numbers and
are never read by a test.
"""

import openpyxl
import pytest

from helpers import GOOD_DATE, GOOD_REMARK, statement_row
from openfloat_formatter.models import StatementReport, StatementTransaction
from openfloat_formatter.normalizer import parse_case_remark
from openfloat_formatter.statement import build_statement_report
from openfloat_formatter.writer import write_finance_workbook, write_statement_workbook

STATEMENT_HEADERS = [
    "Source File",
    "Row",
    "Date",
    "Account Name",
    "Account Number",
    "Status",
    "Reference Id",
    "Remark",
    "Amount",
]


def _load(write, report):
    """Write a report with `write` and open the result."""
    buffer = write(report)
    buffer.seek(0)
    return openpyxl.load_workbook(buffer)


def _rows(worksheet):
    return list(worksheet.iter_rows(min_row=2, values_only=True))


@pytest.fixture
def report(make_statement_workbook):
    """Two paid rows (100 + 250) and one Reversed row."""
    buffer = make_statement_workbook(
        rows=[
            statement_row(),
            statement_row(phone=254798765432, amount=250),
            statement_row(
                status="Reversed",
                phone=254722334455,
                amount=None,
                **{"Reference Id": "REF9"},
            ),
        ]
    )
    return build_statement_report([buffer], source_names=["august.xlsx"])


class TestSheets:
    """Which sheets the workbook has, and what lands on each."""

    def test_two_sheets_without_reconciliation(self, report):
        workbook = _load(write_statement_workbook, report)
        assert workbook.sheetnames == ["Successful", "Unsuccessful"]
        workbook.close()

    def test_reconciliation_sheets_when_an_input_was_supplied(
        self, make_statement_workbook, pm_input_df
    ):
        buffer = make_statement_workbook(rows=[statement_row()])
        reconciled = build_statement_report([buffer], input_df=pm_input_df)
        workbook = _load(write_statement_workbook, reconciled)
        assert workbook.sheetnames == [
            "Successful",
            "Unsuccessful",
            "Paid",
            "Matched but Unpaid",
            "Missing from Statement",
            "Not in Input",
        ]
        workbook.close()

    def test_headers(self, report):
        workbook = _load(write_statement_workbook, report)
        for title in ("Successful", "Unsuccessful"):
            assert [cell.value for cell in workbook[title][1]] == STATEMENT_HEADERS
        workbook.close()

    def test_successful_rows_only_on_the_first_sheet(self, report):
        workbook = _load(write_statement_workbook, report)
        statuses = [row[5] for row in _rows(workbook["Successful"])][:-1]  # drop TOTAL
        assert statuses == ["Successful", "Successful"]
        workbook.close()

    def test_reversed_row_lands_unsuccessful_with_no_amount(self, report):
        workbook = _load(write_statement_workbook, report)
        data = _rows(workbook["Unsuccessful"])[:-1]  # drop TOTAL
        assert len(data) == 1
        row = data[0]
        assert row[5] == "Reversed"
        assert row[6] == "REF9"  # Reference Id survives
        assert row[8] is None  # Amount blank, not zero
        workbook.close()

    def test_source_file_recorded(self, report):
        workbook = _load(write_statement_workbook, report)
        assert _rows(workbook["Successful"])[0][0] == "august.xlsx"
        workbook.close()


class TestTotals:
    """The bold TOTAL row under each sheet."""

    def test_total_sums_the_amount_column(self, report):
        workbook = _load(write_statement_workbook, report)
        worksheet = workbook["Successful"]
        last = _rows(worksheet)[-1]
        assert last[0] == "TOTAL"
        assert last[8] == 350.0  # 100 + 250
        workbook.close()

    def test_total_row_is_bold_and_formatted(self, report):
        workbook = _load(write_statement_workbook, report)
        worksheet = workbook["Successful"]
        label = worksheet.cell(row=worksheet.max_row, column=1)
        amount = worksheet.cell(row=worksheet.max_row, column=9)
        assert label.font.bold
        assert amount.font.bold
        assert amount.number_format == "#,##0"
        workbook.close()

    def test_reversed_rows_total_to_zero(self, report):
        """The Reversed row carries no amount, so its sheet totals nothing."""
        workbook = _load(write_statement_workbook, report)
        assert _rows(workbook["Unsuccessful"])[-1][8] == 0.0
        workbook.close()

    def test_empty_sheet_has_no_total_row(self, make_statement_workbook):
        """Headers only — a bold zero would read like a finding."""
        buffer = make_statement_workbook(rows=[statement_row()])
        workbook = _load(write_statement_workbook, build_statement_report([buffer]))
        worksheet = workbook["Unsuccessful"]
        assert worksheet.max_row == 1
        assert _rows(worksheet) == []
        workbook.close()


class TestCellTypes:
    """Phones as numbers, and free text that cannot become a formula."""

    def test_account_number_written_as_a_number(self, report):
        workbook = _load(write_statement_workbook, report)
        cell = workbook["Successful"].cell(row=2, column=5)
        assert cell.value == 254712345678
        assert isinstance(cell.value, int)
        assert cell.number_format == "0"
        workbook.close()

    def test_free_text_cannot_become_a_formula(self):
        """A Remark that looks like a formula is written as text, not executed.

        Built as a model rather than through a statement file on purpose: reading
        a .xlsx already neutralizes formulas, so this pins the writer's own guard.
        """
        report = StatementReport(
            transactions=[
                StatementTransaction(
                    file_name="august.xlsx",
                    row_number=2,
                    status="Successful",
                    is_successful=True,
                    account_name="TEST001",
                    account_number="254712345678",
                    remark='=HYPERLINK("http://evil","click")',
                    amount=100.0,
                )
            ]
        )
        workbook = _load(write_statement_workbook, report)
        cell = workbook["Successful"].cell(row=2, column=8)
        assert cell.data_type != "f"
        assert str(cell.value).startswith("'=")
        workbook.close()


class TestReconciliationSheets:
    """The follow-up list, which is the part that is painful to copy off a screen."""

    def test_bucket_rows_and_notes(self, make_statement_workbook, pm_input_df):
        buffer = make_statement_workbook(rows=[statement_row()])
        reconciled = build_statement_report([buffer], input_df=pm_input_df)
        workbook = _load(write_statement_workbook, reconciled)

        assert [cell.value for cell in workbook["Paid"][1]] == [
            "Phone",
            "Unique ID",
            "Input Amount",
            "Input Rows",
            "Successful",
            "Unsuccessful",
            "Paid Total",
            "Notes",
        ]
        missing = _rows(workbook["Missing from Statement"])[:-1]
        assert len(missing) > 0
        assert all(isinstance(row[0], int) for row in missing)  # phones as numbers
        workbook.close()

    def test_bucket_totals_the_paid_column(self, make_statement_workbook, pm_input_df):
        buffer = make_statement_workbook(rows=[statement_row()])
        reconciled = build_statement_report([buffer], input_df=pm_input_df)
        workbook = _load(write_statement_workbook, reconciled)
        worksheet = workbook["Paid"]
        total = _rows(worksheet)[-1]
        assert total[0] == "TOTAL"
        assert total[6] == 100.0  # the one successful payment
        workbook.close()


FINANCE_HEADERS = ["Date", "Account Name", "Phone", "Case", "Status", "Debit"]


def _model_txn(status, is_successful, amount, remark=GOOD_REMARK, phone="254712345678"):
    """A transaction built as a model, so an amount survives on a failed row.

    `remark_parts` is parsed the way statement.py parses it, so these behave
    like transactions read from a real export.
    """
    return StatementTransaction(
        file_name="august.xlsx",
        row_number=2,
        status=status,
        is_successful=is_successful,
        date_raw=GOOD_DATE,
        account_name="TEST001",
        account_number=phone,
        remark=remark,
        remark_parts=parse_case_remark(remark)[0],
        amount=amount,
    )


class TestFinanceWorkbook:
    """The sheet finance posts from: one Debit column, failures shaded not dropped."""

    def test_single_sheet_with_ledger_columns(self, report):
        workbook = _load(write_finance_workbook, report)
        assert workbook.sheetnames == ["Finance Reconciliation"]
        assert [cell.value for cell in workbook.active[1]] == FINANCE_HEADERS
        workbook.close()

    def test_successful_row_carries_its_amount_as_debit(self, report):
        workbook = _load(write_finance_workbook, report)
        worksheet = workbook.active
        assert worksheet.cell(row=2, column=6).value == 100
        assert worksheet.cell(row=2, column=5).value == "Successful"
        workbook.close()

    def test_reversed_row_is_shaded_with_no_debit(self, report):
        workbook = _load(write_finance_workbook, report)
        worksheet = workbook.active
        reversed_row = next(
            row for row in range(2, worksheet.max_row + 1)
            if worksheet.cell(row=row, column=5).value == "Reversed"
        )
        assert worksheet.cell(row=reversed_row, column=6).value is None
        assert worksheet.cell(row=reversed_row, column=1).fill.start_color.rgb == "00FFC7CE"
        workbook.close()

    def test_total_is_the_debit_column_only(self, report):
        workbook = _load(write_finance_workbook, report)
        worksheet = workbook.active
        total_row = worksheet.max_row
        assert worksheet.cell(row=total_row, column=1).value == "TOTAL"
        assert worksheet.cell(row=total_row, column=6).value == 350.0  # 100 + 250
        assert worksheet.cell(row=total_row, column=6).font.bold
        assert worksheet.cell(row=total_row, column=6).number_format == "#,##0"
        # "they only want the total for Debit" — nothing else is totalled
        for column in range(2, 6):
            assert worksheet.cell(row=total_row, column=column).value in ("", None)
        workbook.close()

    def test_failed_row_with_an_amount_is_excluded(self):
        """The reason Debit keys off is_successful, not off a missing amount.

        A Reversed row happens to arrive with no amount; a Failed one need not,
        and it is still not money that left the float.
        """
        report = StatementReport(
            transactions=[
                _model_txn("Successful", True, 400.0),
                _model_txn("Failed", False, 600.0, phone="254798765432"),
            ]
        )
        workbook = _load(write_finance_workbook, report)
        worksheet = workbook.active
        assert worksheet.cell(row=3, column=5).value == "Failed"
        assert worksheet.cell(row=3, column=6).value is None
        assert worksheet.cell(row=3, column=1).fill.start_color.rgb == "00FFC7CE"
        assert worksheet.cell(row=worksheet.max_row, column=6).value == 400.0
        workbook.close()

    def test_case_column(self):
        """The parsed case number, or the raw remark when it does not parse."""
        report = StatementReport(
            transactions=[
                _model_txn("Successful", True, 100.0),
                _model_txn("Successful", True, 100.0, remark="no case here"),
            ]
        )
        workbook = _load(write_finance_workbook, report)
        worksheet = workbook.active
        assert worksheet.cell(row=2, column=4).value == "C#37154"
        assert worksheet.cell(row=3, column=4).value == "no case here"
        workbook.close()

    def test_phone_written_as_a_number(self, report):
        workbook = _load(write_finance_workbook, report)
        cell = workbook.active.cell(row=2, column=3)
        assert cell.value == 254712345678
        assert cell.number_format == "0"
        workbook.close()

    def test_no_transactions_means_no_total_row(self):
        workbook = _load(write_finance_workbook, StatementReport())
        assert workbook.active.max_row == 1
        workbook.close()
